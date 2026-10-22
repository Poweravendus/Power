#!/usr/bin/env python3
"""Recalculate dam_annual_analysis.xlsx values (as embedded by the builder's LibreOffice
pass) against an independent pandas computation. Exit code 1 on any mismatch."""
import os, sys
import pandas as pd
from openpyxl import load_workbook
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import annual_dam_analysis as A

HERE = os.path.dirname(os.path.abspath(__file__))
df = pd.read_csv(os.path.join(HERE, "dam_volumes.csv"))
df["norm"] = df.entity.map(A.normalize); df["player"] = df.norm.map(A.player_of); df["group"] = df.player.map(A.group_of)
yr = df.month.str[:4].astype(int); mo = df.month.str[5:7].astype(int); df["cy"] = "CY" + yr.astype(str).str[2:]
fs = yr.where(mo >= 4, yr - 1); df["fy"] = "FY" + fs.astype(str) + "-" + ((fs + 1) % 100).astype(str).str.zfill(2)
wb = load_workbook(sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "dam_annual_analysis.xlsx"), data_only=True)
errs = checked = 0
R = lambda ws: list(ws.iter_rows(values_only=True))
def bad(*m):
    global errs; errs += 1
    if errs <= 20: print("MISMATCH:", *m)
for kind in ["sold", "purchased"]:
    rows = R(wb[f"Monthly_{kind.title()}"])[1:]
    sub = df[df.table == kind].sort_values(["month", "volume_mu"], ascending=[True, False]).reset_index(drop=True)
    assert len(rows) == len(sub)
    for r, (_, e) in zip(rows, sub.iterrows()):
        checked += 1
        if r[1] != e.cy or r[2] != e.fy or r[5] != e.player or r[6] != e.group: bad("monthly", r[:7])
for basis, col in [("CY", "cy"), ("FY", "fy")]:
    for kind in ["sold", "purchased"]:
        for level, pre in [("player", "Annual"), ("group", "AnnualGroup")]:
            ws = wb[f"{pre}_{kind.title()}_{basis}"]; rows = R(ws)
            hdr = rows[1]; klen = 4 if basis == "CY" else 9
            vals = [h[:klen] for h in hdr[1:] if h not in (None, "TOTAL") and not str(h).startswith("Rank")]
            exp = df[df.table == kind].pivot_table(index=level, columns=col, values="volume_mu", aggfunc="sum").fillna(0)
            body = [r for r in rows[2:] if r[0] not in (None, "TOTAL") and not str(r[0]).startswith("Check")]
            for r in body:
                for j, v in enumerate(vals):
                    checked += 1
                    if abs((r[1 + j] or 0) - exp.loc[r[0], v]) > 0.011: bad("annual", ws.title, r[0], v)
            tot = [r for r in rows if r[0] == "TOTAL"][0]; chk = [r for r in rows if str(r[0]).startswith("Check")][0]
            checked += 1
            if abs(tot[1 + len(vals)] - df[df.table == kind].volume_mu.sum()) > 0.011: bad("total", ws.title)
            if any(c != "OK" for c in chk[1:2 + len(vals)]): bad("check row", ws.title)
            for k in range(len(vals) + 1):
                checked += 1
                if sorted(r[1 + len(vals) + 1 + k] for r in body) != list(range(1, len(body) + 1)): bad("rank not unique", ws.title, hdr[1 + len(vals) + 1 + k])
            tws = wb[f"{'Top15' if level == 'player' else 'Top15Group'}_{kind.title()}_{basis}"]; trows = R(tws)
            expt = exp.copy(); expt["TOTAL"] = expt.sum(axis=1); expt = expt.sort_values("TOTAL", ascending=False)
            top = trows[2:17]
            if [r[1] for r in top] != list(expt.index[:15]): bad("top15 order", tws.title)
            for r in top:
                for j, v in enumerate(vals + ["TOTAL"]):
                    checked += 1
                    if abs((r[2 + j] or 0) - expt.loc[r[1], v]) > 0.011: bad("top15 value", tws.title, r[1], v)
            others = trows[17]; chk = trows[19]
            for j, v in enumerate(vals + ["TOTAL"]):
                checked += 1
                if abs(others[2 + j] - expt[v].iloc[15:].sum()) > 0.011: bad("others", tws.title, v)
            if any(c != "OK" for c in chk[2:3 + len(vals)]): bad("top15 check", tws.title)
            sh = trows[23:39]   # 15 players + Others (TOTAL row excluded)
            for j in range(len(vals) + 1):
                checked += 1
                if abs(sum((r[2 + j] or 0) for r in sh) - 100) > 0.02: bad("share sum", tws.title, j)
            if level == "player":
                bws = wb[f"Top15_ByYear_{kind.title()}_{basis}"]; brows = R(bws)
                for j, v in enumerate(vals):
                    c0 = j * 5; s = exp[v].sort_values(ascending=False); s = s[s > 0]
                    gv = [brows[3 + i][c0 + 2] for i in range(15)]; checked += 1
                    if [round(x, 2) for x in gv] != [round(x, 2) for x in s.iloc[:15]]: bad("byyear", bws.title, v)
                    if abs(brows[18][c0 + 2] - s.iloc[15:].sum()) > 0.011 or abs(brows[19][c0 + 2] - s.sum()) > 0.011: bad("byyear others/total", bws.title, v)
                    if abs(brows[19][c0 + 3] - 100) > 0.02: bad("byyear share", bws.title, v)
print(f"checked {checked} cells/conditions, mismatches: {errs}")
sys.exit(1 if errs else 0)
