#!/usr/bin/env python3
"""Formula-driven DAM analysis workbook (dam_annual_analysis.xlsx).

Only the Monthly_* sheets and the Mapping sheet hold hard data. Everything else
is an Excel formula chained off them:

  Monthly_Sold / Monthly_Purchased : one row per (month, printed entity); the
        CY, FY, player and group columns are formulas (year parsed from the month,
        player/group looked up from Mapping).
  Mapping            : printed name -> player -> group. Edit here and every sheet updates.
  Annual_<kind>_<basis>       : SUMIFS of the monthly sheet by player and year, with
        TOTAL, a unique rank on TOTAL, a unique rank per year, and a reconciliation
        row that compares the sheet total with the monthly sheet (must read OK).
  Top15_<kind>_<basis>        : INDEX/MATCH on the rank column of the Annual sheet,
        Others = Total - top 15, plus a share-of-total block.
  Top15_ByYear_<kind>_<basis> : each year's own top 15 from the per-year rank columns.
  AnnualGroup_* / Top15Group_*: the same at parent-group level.

basis: CY = calendar year, FY = April-March fiscal year.
"""
import os, sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter as L

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from annual_dam_analysis import normalize, player_of, group_of   # same mapping rules

HERE = os.path.dirname(os.path.abspath(__file__))
df = pd.read_csv(os.path.join(HERE, "dam_volumes.csv"))
df["norm"] = df["entity"].map(normalize)
df["player"] = df["norm"].map(player_of)
df["group"] = df["player"].map(group_of)
yr, mo = df["month"].str[:4].astype(int), df["month"].str[5:7].astype(int)
df["cy"] = yr
fs = yr.where(mo >= 4, yr - 1)
df["fy"] = "FY" + fs.astype(str) + "-" + ((fs + 1) % 100).astype(str).str.zfill(2)
first, last = df["month"].min(), df["month"].max()

CY = sorted(df["cy"].unique())
FY = sorted(df["fy"].unique())
PERIODS = {"CY": ("cy", CY), "FY": ("fy", FY)}
partial = {}
for basis, (col, vals) in PERIODS.items():
    for v in vals:
        m = sorted(df.loc[df[col] == v, "month"].unique())
        if len(m) < 12:
            partial[(basis, v)] = f"{m[0]} to {m[-1]} only"

BOLD = Font(bold=True)
HEAD = PatternFill("solid", fgColor="DDEBF7")
TOTF = PatternFill("solid", fgColor="FFF2CC")
NUM = "#,##0.00"
PCT = "0.00"

wb = Workbook()
wb.remove(wb.active)


def header(ws, row, values, fill=HEAD):
    for j, v in enumerate(values, 1):
        c = ws.cell(row=row, column=j, value=v)
        c.font = BOLD
        c.fill = fill
        c.alignment = Alignment(horizontal="center", wrap_text=True)


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


# ---------------------------------------------------------------- README
ws = wb.create_sheet("README")
readme = f"""FORMULA-DRIVEN DAM ANALYSIS - sold and purchased volumes in the Day Ahead Market through power exchanges
Source: CERC Market Monitoring monthly reports, {first} to {last}. Units: MU.

HOW THE SHEETS CHAIN (no pasted numbers beyond the first two)
1. Monthly_Sold, Monthly_Purchased  - raw rows as printed in the reports (month, entity, volume, %). Columns CY, FY, Player, Group are formulas.
2. Mapping                          - printed name -> Player -> Group. The only place names are decided. Edit Player/Group here and every downstream sheet recalculates.
3. Annual_*                         - SUMIFS over the Monthly sheet by Player and year. Includes TOTAL, unique rank on TOTAL, unique rank per year, and a reconciliation row (OK / MISMATCH) against the Monthly sheet.
4. Top15_*                          - pulls players ranked 1-15 from Annual_* by INDEX/MATCH; Others = Total - top 15. Share block below.
5. Top15_ByYear_*                   - each year's own top 15 from the per-year rank columns of Annual_*.
6. AnnualGroup_* / Top15Group_*     - same as 3-4 but by parent Group (Adani Green SPVs, NTPC stations etc. rolled up).

PERIODS
- CY = calendar year. FY = April-March fiscal year.
- Partial periods: {"; ".join(f"{b} {v}: {t}" for (b, v), t in partial.items())}.

CHECKS BUILT IN
- Every Annual sheet has a 'Check vs Monthly' row: sheet total minus SUMIFS of the monthly sheet for that year. Shows OK when zero.
- Every Top15 sheet has a check that top 15 + Others equals the Annual total.
- Sold and Purchased totals are equal in every month by construction of the market (buy = sell); compare the TOTAL rows of Annual_Sold and Annual_Purchased.

NAME MAPPING NOTES
- CERC changed naming conventions several times (notably Aug-2024). Renames merged only where the old label stops exactly when the new one starts (e.g. MP -> Madhya Pradesh; RPREL -> ARTPP -> Adani Power Raipur TPP). 'renamed' column in Mapping flags each.
- 'Government of Himachal Pradesh_<plant>' rows (state free-power shares shown separately from Aug-2024) are mapped into HIMACHAL PRADESH; before Aug-2024 they sat inside the single state row.
- Inferred from timing only: AMPS HAZIRA = AM/NS India (Hazira); RENEW_SR_SLR = ReNew Surya Ravi; ADANI_4RE = Adani Renewable Energy Holding Four.
- Plant-level rows introduced from Aug-2024 (Adani Green SPVs, NTPC stations, ReNew SPVs) are NOT merged at Player level; the Group sheets roll them up.
- Workbook is set to recalculate fully on open. If a sheet shows zeros, press Ctrl+Alt+F9 in Excel.
"""
for i, line in enumerate(readme.strip().split("\n"), 1):
    ws.cell(row=i, column=1, value=line)
ws.column_dimensions["A"].width = 160

# ---------------------------------------------------------------- Mapping
mp = (df.groupby("norm").agg(player=("player", "first"), group=("group", "first"),
                             tables=("table", lambda s: "/".join(sorted(set(s)))),
                             first_month=("month", "min"), last_month=("month", "max"),
                             months=("month", "count"), total_mu=("volume_mu", "sum"))
        .reset_index().sort_values("total_mu", ascending=False))
ws = wb.create_sheet("Mapping")
header(ws, 1, ["printed_name (normalised)", "Player  <- edit", "Group  <- edit", "renamed", "tables",
               "first_month", "last_month", "months", "total_mu"])
for i, r in enumerate(mp.itertuples(index=False), 2):
    ws.cell(row=i, column=1, value=r.norm)
    ws.cell(row=i, column=2, value=r.player)
    ws.cell(row=i, column=3, value=r.group)
    ws.cell(row=i, column=4, value="yes" if r.norm != r.player else "")
    ws.cell(row=i, column=5, value=r.tables)
    ws.cell(row=i, column=6, value=r.first_month)
    ws.cell(row=i, column=7, value=r.last_month)
    ws.cell(row=i, column=8, value=int(r.months))
    ws.cell(row=i, column=9, value=round(float(r.total_mu), 2)).number_format = NUM
MAP_N = len(mp) + 1
ws.freeze_panes = "A2"
widths(ws, {"A": 60, "B": 48, "C": 36, "D": 9, "E": 14, "F": 11, "G": 11, "H": 8, "I": 12})

# ---------------------------------------------------------------- Monthly sheets
MON = {}   # kind -> (sheet name, last row)
for kind in ("sold", "purchased"):
    sub = df[df["table"] == kind].sort_values(["month", "volume_mu"], ascending=[True, False])
    name = f"Monthly_{kind.title()}"
    ws = wb.create_sheet(name)
    header(ws, 1, ["month", "CY", "FY", "entity (as printed)", "printed_name (normalised)",
                   "Player", "Group", "volume_mu", "pct_of_table"])
    for i, r in enumerate(sub.itertuples(index=False), 2):
        ws.cell(row=i, column=1, value=r.month)
        ws.cell(row=i, column=2, value=f"=VALUE(LEFT(A{i},4))")
        ws.cell(row=i, column=3, value=f'="FY"&IF(VALUE(MID(A{i},6,2))>=4,VALUE(LEFT(A{i},4)),VALUE(LEFT(A{i},4))-1)'
                                       f'&"-"&TEXT(MOD(IF(VALUE(MID(A{i},6,2))>=4,VALUE(LEFT(A{i},4)),VALUE(LEFT(A{i},4))-1)+1,100),"00")')
        ws.cell(row=i, column=4, value=r.entity)
        ws.cell(row=i, column=5, value=r.norm)
        ws.cell(row=i, column=6, value=f"=INDEX(Mapping!$B$2:$B${MAP_N},MATCH(E{i},Mapping!$A$2:$A${MAP_N},0))")
        ws.cell(row=i, column=7, value=f"=INDEX(Mapping!$C$2:$C${MAP_N},MATCH(E{i},Mapping!$A$2:$A${MAP_N},0))")
        ws.cell(row=i, column=8, value=float(r.volume_mu)).number_format = NUM
        ws.cell(row=i, column=9, value=None if pd.isna(r.pct_of_volume) else float(r.pct_of_volume)).number_format = PCT
    n = len(sub) + 1
    MON[kind] = (name, n)
    ws.freeze_panes = "A2"
    widths(ws, {"A": 10, "B": 7, "C": 11, "D": 55, "E": 55, "F": 45, "G": 34, "H": 12, "I": 10})
    ws.auto_filter.ref = f"A1:I{n}"


# ---------------------------------------------------------------- Annual + Top15 builders
def annual_sheet(sheet, kind, basis, level):
    """level: 'player' (col F) or 'group' (col G) of the monthly sheet."""
    mname, mn = MON[kind]
    pcol, vals = PERIODS[basis]
    keycol = "F" if level == "player" else "G"
    ycol = "B" if basis == "CY" else "C"
    keys = sorted(df.loc[df["table"] == kind, level].unique())
    ws = wb.create_sheet(sheet)
    ny = len(vals)
    # columns: A key | B.. years | TOTAL | Rank(TOTAL) | rank per year...
    cT = 2 + ny                      # TOTAL col index
    cR = cT + 1                      # rank on total
    cY0 = cR + 1                     # first per-year rank col
    ws.cell(row=1, column=1, value=f"{kind.upper()} in DAM by {level} and {basis} (MU). Formulas: SUMIFS over {mname}.").font = BOLD
    for j, v in enumerate(vals):
        if (basis, v) in partial:
            ws.cell(row=1, column=2 + j, value=partial[(basis, v)]).font = Font(italic=True, color="C00000")
    header(ws, 2, [level.title()] + list(vals) + ["TOTAL", "Rank (TOTAL)"] + [f"Rank {v}" for v in vals])
    r0, r1 = 3, 2 + len(keys)
    for i, k in enumerate(keys, r0):
        ws.cell(row=i, column=1, value=k)
        for j, v in enumerate(vals):
            c = L(2 + j)
            crit = v if isinstance(v, (int, float)) else f'"{v}"'
            ws.cell(row=i, column=2 + j,
                    value=f"=ROUND(SUMIFS({mname}!$H$2:$H${mn},{mname}!${keycol}$2:${keycol}${mn},$A{i},{mname}!${ycol}$2:${ycol}${mn},{c}$2),4)"
                    ).number_format = NUM
        t = L(cT)
        ws.cell(row=i, column=cT, value=f"=ROUND(SUM({L(2)}{i}:{L(1 + ny)}{i}),4)").number_format = NUM
        ws.cell(row=i, column=cR, value=f"=RANK({t}{i},${t}${r0}:${t}${r1},0)+COUNTIF(${t}${r0}:{t}{i},{t}{i})-1")
        for j in range(ny):
            c = L(2 + j)
            ws.cell(row=i, column=cY0 + j,
                    value=f"=RANK({c}{i},${c}${r0}:${c}${r1},0)+COUNTIF(${c}${r0}:{c}{i},{c}{i})-1")
    rt = r1 + 1
    ws.cell(row=rt, column=1, value="TOTAL").font = BOLD
    for j in range(ny + 1):
        c = L(2 + j)
        cell = ws.cell(row=rt, column=2 + j, value=f"=SUM({c}{r0}:{c}{r1})")
        cell.number_format = NUM; cell.font = BOLD; cell.fill = TOTF
    rc = rt + 1
    ws.cell(row=rc, column=1, value="Check vs Monthly (should be OK)").font = BOLD
    for j, v in enumerate(vals):
        c = L(2 + j)
        ws.cell(row=rc, column=2 + j,
                value=f'=IF(ABS({c}{rt}-SUMIFS({mname}!$H$2:$H${mn},{mname}!${ycol}$2:${ycol}${mn},{c}$2))<0.01,"OK","MISMATCH")')
    t = L(cT)
    ws.cell(row=rc, column=cT, value=f'=IF(ABS({t}{rt}-SUM({mname}!$H$2:$H${mn}))<0.01,"OK","MISMATCH")')
    ws.freeze_panes = "B3"
    widths(ws, {"A": 50})
    for j in range(2, cY0 + ny):
        ws.column_dimensions[L(j)].width = 13
    return dict(sheet=sheet, r0=r0, r1=r1, rt=rt, cT=cT, cR=cR, cY0=cY0, ny=ny, vals=vals, nkeys=len(keys))


def top15_sheet(sheet, kind, a, level, n=15):
    A, vals, ny = a["sheet"], a["vals"], a["ny"]
    ws = wb.create_sheet(sheet)
    ws.cell(row=1, column=1, value=f"{kind.upper()} in DAM - top {n} {level}s by TOTAL volume over the whole period (MU), from {A}").font = BOLD
    header(ws, 2, ["Rank", level.title()] + list(vals) + ["TOTAL"])
    keyrng = f"{A}!$A${a['r0']}:$A${a['r1']}"
    rankrng = f"{A}!${L(a['cR'])}${a['r0']}:${L(a['cR'])}${a['r1']}"
    for i in range(n):
        r = 3 + i
        ws.cell(row=r, column=1, value=i + 1)
        ws.cell(row=r, column=2, value=f"=INDEX({keyrng},MATCH(A{r},{rankrng},0))")
        for j in range(ny + 1):
            c = L(2 + j)
            ws.cell(row=r, column=3 + j,
                    value=f"=INDEX({A}!${c}${a['r0']}:${c}${a['r1']},MATCH($B{r},{keyrng},0))").number_format = NUM
    ro, rt, rc = 3 + n, 4 + n, 5 + n
    ws.cell(row=ro, column=1, value=f'=COUNTIF({A}!${L(a["cT"])}${a["r0"]}:${L(a["cT"])}${a["r1"]},">0")-{n}')
    ws.cell(row=ro, column=2, value=f'="OTHERS ("&A{ro}&" {level}s)"').font = BOLD
    ws.cell(row=rt, column=2, value="TOTAL").font = BOLD
    ws.cell(row=rc, column=2, value="Check: top15 + Others = Annual total").font = BOLD
    for j in range(ny + 1):
        c = L(3 + j)
        ac = L(2 + j)
        ws.cell(row=ro, column=3 + j, value=f"={c}{rt}-SUM({c}3:{c}{2 + n})").number_format = NUM
        tc = ws.cell(row=rt, column=3 + j, value=f"={A}!{ac}{a['rt']}")
        tc.number_format = NUM; tc.font = BOLD; tc.fill = TOTF
        ws.cell(row=rc, column=3 + j, value=f'=IF(ABS(SUM({c}3:{c}{ro})-{c}{rt})<0.01,"OK","MISMATCH")')
    # share block
    rs = rc + 2
    ws.cell(row=rs, column=1, value="Share of period total (%)").font = BOLD
    header(ws, rs + 1, ["Rank", level.title()] + list(vals) + ["TOTAL"])
    for i in range(n + 2):
        src = 3 + i
        r = rs + 2 + i
        ws.cell(row=r, column=1, value=f"=A{src}")
        ws.cell(row=r, column=2, value=f"=B{src}")
        for j in range(ny + 1):
            c = L(3 + j)
            ws.cell(row=r, column=3 + j, value=f"=IF({c}${rt}=0,0,{c}{src}/{c}${rt}*100)").number_format = PCT
    for j, v in enumerate(vals):
        if (a["sheet"].endswith("CY") or a["sheet"].endswith("FY")) and (("CY" if a["sheet"].endswith("CY") else "FY"), v) in partial:
            ws.cell(row=1, column=3 + j, value=partial[(("CY" if a["sheet"].endswith("CY") else "FY"), v)]).font = Font(italic=True, color="C00000")
    ws.freeze_panes = "C3"
    widths(ws, {"A": 7, "B": 50})
    for j in range(ny + 1):
        ws.column_dimensions[L(3 + j)].width = 14


def top15_by_year_sheet(sheet, kind, a, n=15):
    A, vals, ny = a["sheet"], a["vals"], a["ny"]
    ws = wb.create_sheet(sheet)
    ws.cell(row=1, column=1, value=f"{kind.upper()} in DAM - each period's own top {n} players (MU and % of that period), from {A}").font = BOLD
    keyrng = f"{A}!$A${a['r0']}:$A${a['r1']}"
    for j, v in enumerate(vals):
        c0 = 1 + j * 5                       # block start column
        vcol = L(2 + j)                       # value column in Annual
        rcol = L(a["cY0"] + j)                # per-year rank column in Annual
        lab = f"{v}" + (f"  ({partial[('CY' if A.endswith('CY') else 'FY', v)]})" if ('CY' if A.endswith('CY') else 'FY', v) in partial else "")
        ws.cell(row=2, column=c0, value=lab).font = BOLD
        header(ws, 3, [None] * (c0 - 1) + ["Rank", "Player", "MU", "% share"])
        for i in range(n):
            r = 4 + i
            ws.cell(row=r, column=c0, value=i + 1)
            ws.cell(row=r, column=c0 + 1, value=f"=INDEX({keyrng},MATCH({L(c0)}{r},{A}!${rcol}${a['r0']}:${rcol}${a['r1']},0))")
            ws.cell(row=r, column=c0 + 2, value=f"=INDEX({A}!${vcol}${a['r0']}:${vcol}${a['r1']},MATCH({L(c0 + 1)}{r},{keyrng},0))").number_format = NUM
            ws.cell(row=r, column=c0 + 3, value=f"=IF({L(c0 + 2)}${4 + n + 1}=0,0,{L(c0 + 2)}{r}/{L(c0 + 2)}${4 + n + 1}*100)").number_format = PCT
        ro, rt = 4 + n, 5 + n
        ws.cell(row=ro, column=c0 + 1, value=f'="OTHERS ("&(COUNTIF({A}!${vcol}${a["r0"]}:${vcol}${a["r1"]},">0")-{n})&" players)"').font = BOLD
        ws.cell(row=ro, column=c0 + 2, value=f"={L(c0 + 2)}{rt}-SUM({L(c0 + 2)}4:{L(c0 + 2)}{3 + n})").number_format = NUM
        ws.cell(row=ro, column=c0 + 3, value=f"=IF({L(c0 + 2)}{rt}=0,0,{L(c0 + 2)}{ro}/{L(c0 + 2)}{rt}*100)").number_format = PCT
        ws.cell(row=rt, column=c0 + 1, value="TOTAL").font = BOLD
        tc = ws.cell(row=rt, column=c0 + 2, value=f"={A}!{vcol}{a['rt']}")
        tc.number_format = NUM; tc.font = BOLD; tc.fill = TOTF
        ws.cell(row=rt, column=c0 + 3, value=f"=IF({L(c0 + 2)}{rt}=0,0,SUM({L(c0 + 3)}4:{L(c0 + 3)}{ro}))").number_format = PCT
        ws.column_dimensions[L(c0)].width = 6
        ws.column_dimensions[L(c0 + 1)].width = 44
        ws.column_dimensions[L(c0 + 2)].width = 12
        ws.column_dimensions[L(c0 + 3)].width = 9
        ws.column_dimensions[L(c0 + 4)].width = 3
    ws.freeze_panes = "A4"


# header row 3 in by-year sheet: fix header() call which pads with None (skip None cells)
_orig_header = header
def header(ws, row, values, fill=HEAD):   # noqa: F811
    for j, v in enumerate(values, 1):
        if v is None:
            continue
        c = ws.cell(row=row, column=j, value=v)
        c.font = BOLD; c.fill = fill
        c.alignment = Alignment(horizontal="center", wrap_text=True)


for basis in ("CY", "FY"):
    for kind in ("sold", "purchased"):
        K = kind.title()
        a = annual_sheet(f"Annual_{K}_{basis}", kind, basis, "player")
        top15_sheet(f"Top15_{K}_{basis}", kind, a, "player")
        top15_by_year_sheet(f"Top15_ByYear_{K}_{basis}", kind, a)
        g = annual_sheet(f"AnnualGroup_{K}_{basis}", kind, basis, "group")
        top15_sheet(f"Top15Group_{K}_{basis}", kind, g, "group")

# validation sheet from extraction step (values)
val = pd.read_csv(os.path.join(HERE, "dam_volumes_validation.csv"))
ws = wb.create_sheet("PDF_Extraction_Check")
header(ws, 1, list(val.columns))
for i, r in enumerate(val.itertuples(index=False), 2):
    for j, v in enumerate(r, 1):
        ws.cell(row=i, column=j, value=None if pd.isna(v) else (float(v) if isinstance(v, float) else v))
widths(ws, {"A": 10, "B": 11, "C": 12, "D": 7, "E": 13, "F": 14, "G": 10})

# order: README, Mapping, Monthly x2, then CY sheets, FY sheets, check
order = ["README", "Mapping", "Monthly_Sold", "Monthly_Purchased"]
order += [s.title for s in wb.worksheets if s.title.endswith("_CY")]
order += [s.title for s in wb.worksheets if s.title.endswith("_FY")]
order += ["PDF_Extraction_Check"]
wb._sheets = [wb[s] for s in order]
wb.calculation.fullCalcOnLoad = True
out = os.path.join(HERE, "dam_annual_analysis.xlsx")
wb.save(out)
print("wrote", out, "sheets:", len(wb.worksheets))

# ---------------------------------------------------------------- recalculate (optional)
# openpyxl stores formulas without cached results; viewers that do not recalculate would show
# blanks. If LibreOffice is available, open+recalculate+resave so the file carries both.
import shutil, subprocess, tempfile
if shutil.which("soffice"):
    tmp = tempfile.mkdtemp()
    prof = os.path.join(tmp, "profile", "user")
    os.makedirs(prof)
    with open(os.path.join(prof, "registrymodifications.xcu"), "w") as fh:
        fh.write('<?xml version="1.0" encoding="UTF-8"?><oor:items xmlns:oor="http://openoffice.org/2001/registry" '
                 'xmlns:xs="http://www.w3.org/2001/XMLSchema" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
                 '<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse">'
                 '<value>0</value></prop></item></oor:items>')
    r = subprocess.run(["soffice", "--headless", "--norestore", f"-env:UserInstallation=file://{tmp}/profile",
                        "--convert-to", "xlsx", "--outdir", os.path.join(tmp, "out"), out],
                       capture_output=True, text=True)
    recalced = os.path.join(tmp, "out", os.path.basename(out))
    if os.path.exists(recalced):
        shutil.copy(recalced, out)
        print("recalculated with LibreOffice; cached values embedded")
    else:
        print("LibreOffice recalculation failed; file has formulas only:", r.stderr[-300:])
    shutil.rmtree(tmp, ignore_errors=True)
