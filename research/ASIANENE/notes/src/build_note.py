"""Markdown -> Word in the Arrow Greentech note format.
Arial; 9pt body; 11pt bold section heads; 18pt title; Letter page, 0.9in margins; grey (D9D9D9) table header fills with borders;
full-width charts via lines of the form  ![caption](path.png) ; '::: small' lines render 8pt grey source notes.
Usage: python3 build_note.py source.md out.docx "Header text" """
import re, sys
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC, OUT, HDR = sys.argv[1], sys.argv[2], sys.argv[3]
GREY = RGBColor(0x59, 0x59, 0x59); RED = RGBColor(0xEE, 0x00, 0x00); INK = RGBColor(0, 0, 0)
doc = Document()
sec = doc.sections[0]; sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Inches(0.9)
USABLE = Inches(8.5 - 1.8)
N = doc.styles['Normal']; N.font.name = 'Arial'; N.font.size = Pt(9); N.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
N.paragraph_format.space_after = Pt(4); N.paragraph_format.line_spacing = 1.08

def add_page_field(par):
    r = par.add_run(); f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin')
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = 'PAGE'
    f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'end'); r._r.append(f1); r._r.append(it); r._r.append(f2)
    r.font.size = Pt(8); r.font.color.rgb = GREY
hp = sec.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT; hr = hp.add_run(HDR); hr.font.size = Pt(8); hr.font.color.rgb = GREY
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = fp.add_run('Strictly Confidential - for internal circulation only. This note is for informational purposes and does not constitute investment advice. Page ')
fr.font.size = Pt(8); fr.font.color.rgb = GREY; add_page_field(fp)

INLINE = re.compile(r'(\*\*[^*]+\*\*|\*[^*]+\*)')
def add_inline(par, text, size=None, bold=False, color=None, italic=False):
    for tok in INLINE.split(text):
        if not tok: continue
        if tok.startswith('**') and tok.endswith('**'): r = par.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith('*') and tok.endswith('*') and len(tok) > 2: r = par.add_run(tok[1:-1]); r.italic = True
        else: r = par.add_run(tok)
        if size: r.font.size = Pt(size)
        if bold: r.bold = True
        if italic: r.italic = True
        if color: r.font.color.rgb = color

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def add_table(rows):
    ncol = len(rows[0]); import math
    weights = []
    for j in range(ncol):
        lens = [len(r[j]) for r in rows]; weights.append(max(1.8, (sum(lens) / len(lens)) ** 0.72))
    tot = sum(weights); widths = [max(int(Cm(1.2)), int(USABLE * w / tot)) for w in weights]
    over = sum(widths) - USABLE
    if over > 0:
        big = [k for k in range(ncol) if widths[k] > int(Cm(1.2))]; s = sum(widths[k] for k in big)
        widths = [int(w - over * (w / s)) if k in big else w for k, w in enumerate(widths)]
    fs = 8 if ncol <= 6 else 7.5
    t = doc.add_table(rows=len(rows), cols=ncol); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    tblPr = t._tbl.tblPr; lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    mar = OxmlElement('w:tblCellMar')
    for side, v in (('top', 30), ('left', 60), ('bottom', 30), ('right', 60)):
        e = OxmlElement(f'w:{side}'); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); mar.append(e)
    tblPr.append(mar)
    for j in range(ncol): t.columns[j].width = widths[j]
    for i, row in enumerate(rows):
        tr = t.rows[i]
        if i == 0:
            trPr = tr._tr.get_or_add_trPr(); h = OxmlElement('w:tblHeader'); h.set(qn('w:val'), 'true'); trPr.append(h)
        for j in range(ncol):
            c = tr.cells[j]; c.width = widths[j]; txt = row[j] if j < len(row) else ''
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
            if i > 0 and j > 0 and re.fullmatch(r'[\s\d.,%()x+\-–−/]*\d[\s\d.,%()x+\-–−/]*', txt or ''): p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            add_inline(p, txt, size=fs, bold=(i == 0))
            if i == 0: shade(c, 'D9D9D9')
    sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(2)

def flush_table(buf):
    rows = []
    for l in buf:
        cells = [c.strip() for c in l.strip().strip('|').split('|')]
        if all(re.fullmatch(r':?-{2,}:?', c) for c in cells if c): continue
        rows.append(cells)
    if not rows: return
    ncol = max(len(r) for r in rows); add_table([r + [''] * (ncol - len(r)) for r in rows])

lines = open(SRC, encoding='utf-8').read().split('\n'); tbuf = []; first_h1 = True
for raw in lines:
    l = raw.rstrip()
    if l.startswith('|'): tbuf.append(l); continue
    if tbuf: flush_table(tbuf); tbuf = []
    if not l.strip(): continue
    m = re.match(r'^!\[(.*?)\]\((.*?)\)$', l.strip())
    if m:
        doc.add_picture(m.group(2), width=USABLE)
        if m.group(1):
            p = doc.add_paragraph(); add_inline(p, m.group(1), size=8, color=GREY); p.paragraph_format.space_after = Pt(6)
        continue
    if l.startswith('# '):
        p = doc.add_paragraph(); add_inline(p, l[2:], size=18, bold=True); p.paragraph_format.space_after = Pt(2); continue
    if l.startswith('## '):
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(4); p.paragraph_format.keep_with_next = True
        add_inline(p, l[3:], size=11, bold=True); continue
    if l.startswith('### '):
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(2); p.paragraph_format.keep_with_next = True
        add_inline(p, l[4:], size=9, bold=True); continue
    if l.startswith('::: key'):
        p = doc.add_paragraph(); add_inline(p, l[8:].strip(), size=9, bold=True); p.paragraph_format.space_after = Pt(2); continue
    if l.startswith('::: date'):
        p = doc.add_paragraph(); add_inline(p, l[9:].strip(), size=9, color=GREY); p.paragraph_format.space_after = Pt(8); continue
    if l.startswith('::: small'):
        p = doc.add_paragraph(); add_inline(p, l[10:].strip(), size=8, color=GREY); p.paragraph_format.space_after = Pt(6); continue
    if l.startswith('::: red'):
        p = doc.add_paragraph(); add_inline(p, l[8:].strip(), size=9, bold=True, color=RED); continue
    if l.strip() == '---':
        p = doc.add_paragraph(); pPr = p._p.get_or_add_pPr(); bdr = OxmlElement('w:pBdr'); b = OxmlElement('w:bottom')
        b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '6'); b.set(qn('w:space'), '1'); b.set(qn('w:color'), '999999'); bdr.append(b); pPr.append(bdr); continue
    if re.match(r'^\s*[-*] ', l):
        indent = (len(l) - len(l.lstrip())) // 2
        p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.left_indent = Cm(0.55 + 0.5 * indent); p.paragraph_format.space_after = Pt(3)
        add_inline(p, re.sub(r'^\s*[-*] ', '', l)); continue
    if re.match(r'^\s*\d+\. ', l):
        m = re.match(r'^(\s*)(\d+)\. (.*)$', l)
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.8); p.paragraph_format.first_line_indent = Cm(-0.5); p.paragraph_format.space_after = Pt(3)
        add_inline(p, f'{m.group(2)}.  {m.group(3)}'); continue
    p = doc.add_paragraph(); add_inline(p, l.strip())
if tbuf: flush_table(tbuf)
doc.save(OUT); print('saved', OUT)
