#!/usr/bin/env python3
"""Pivot the shareholding database into a readable % matrix: categories down, quarter-ends across."""
from pathlib import Path
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parent / "output"
BUCKETS = [("promoter", "Promoter & promoter group"), ("fii", "FII / FPI"), ("dii", "DII"),
           ("government", "Government"), ("kmp", "Key Managerial Personnel"),
           ("public", "Public (non-institutional, excl. KMP)"), ("others", "Others (employee trust)")]
# SEBI sub-categories worth showing under the buckets (category_key -> label)
SUB = [
    ("fii", "institutionsforeignportfolioinvestorcategoryone", "  FPI category I"),
    ("fii", "institutionsforeignportfolioinvestorcategorytwo", "  FPI category II"),
    ("fii", "institutionsforeignportfolioinvestor", "  FPI (old format, undivided)"),
    ("dii", "mutualfundsoruti", "  Mutual funds / UTI"),
    ("dii", "insurancecompanies", "  Insurance companies"),
    ("dii", "alternativeinvestmentfunds", "  Alternative investment funds"),
    ("dii", "banks", "  Banks"),
    ("dii", "financialinstitutionorbanks", "  Financial institutions / banks (old format)"),
    ("dii", "nbfcsregisteredwithrbi", "  NBFCs"),
    ("dii", "otherinstitutions", "  Other institutions incl. insurance (old format)"),
    ("public", "residentindividualshareholdersholdingnominalsharecapitaluptorstwolakh", "  Resident individuals up to Rs 2 lakh"),
    ("public", "individualshareholdersholdingnominalsharecapitaluptorstwolakh", "  Individuals up to Rs 2 lakh (old format)"),
    ("public", "residentindividualshareholdersholdingnominalsharecapitalinexcessofrstwolakh", "  Resident individuals above Rs 2 lakh"),
    ("public", "individualshareholdersholdingnominalsharecapitalinexcessofrstwolakh", "  Individuals above Rs 2 lakh (old format)"),
    ("public", "bodiescorporate", "  Bodies corporate (incl. Dalmia group)"),
    ("public", "nonresidentindians", "  Non-resident Indians"),
    ("public", "directorsanddirectorsrelatives", "  Directors & relatives"),
    ("public", "othernoninstitutions", "  Other non-institutions"),
]

s = pd.read_csv(OUT / "iex_shareholding_summary.csv", parse_dates=["quarter_end"])
c = pd.read_csv(OUT / "iex_shareholding_categories_long.csv", parse_dates=["quarter_end"])
s = s[s.quarter_end.dt.is_quarter_end].sort_values("quarter_end")
cols = [d.strftime("%b-%y") for d in s.quarter_end]

rows = []
for key, label in BUCKETS:
    rows.append([label] + [round(v, 2) if pd.notna(v) else None for v in s[f"{key}_pct"]])
    for bkey, ckey, sublabel in SUB:
        if bkey != key:
            continue
        sub = c[c.category_key == ckey].set_index("quarter_end")["pct_of_total"].reindex(s.quarter_end)
        if sub.notna().any() and (sub.fillna(0) != 0).any():
            rows.append([sublabel] + [round(v, 2) if pd.notna(v) else None for v in sub])
rows.append(["Total"] + [round(sum(s[f"{k}_pct"].fillna(0).iloc[i] for k, _ in BUCKETS), 2) for i in range(len(cols))])
shares = ["Total shares (crore)"] + [round(v / 1e7, 2) for v in s.total_shares]

matrix = pd.DataFrame(rows, columns=["Category (% of total shares)"] + cols)
matrix.to_csv(OUT / "iex_shareholding_pct_matrix.csv", index=False)

path = OUT / "iex_shareholding_pct_matrix.xlsx"
with pd.ExcelWriter(path, engine="openpyxl") as xw:
    matrix.to_excel(xw, sheet_name="Shareholding %", index=False, startrow=1)
    ws = xw.book["Shareholding %"]
    ws["A1"] = "Indian Energy Exchange Ltd - shareholding pattern, % of total shares (source: SEBI Reg. 31 filings via NSE)"
    ws["A1"].font = Font(bold=True, size=12)
    hdr_fill = PatternFill("solid", fgColor="1F3864"); bucket_fill = PatternFill("solid", fgColor="DDEBF7")
    thin = Side(style="thin", color="BFBFBF")
    for cell in ws[2]:
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in range(3, 3 + len(rows)):
        label = ws.cell(r, 1).value
        is_bucket = not str(label).startswith("  ")
        for col in range(1, len(cols) + 2):
            cell = ws.cell(r, col)
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            if col > 1:
                cell.number_format = '0.00'; cell.alignment = Alignment(horizontal="right")
            if is_bucket:
                cell.font = Font(bold=True); cell.fill = bucket_fill
            if label == "Total":
                cell.font = Font(bold=True)
    r = 3 + len(rows)
    for i, v in enumerate(shares):
        cell = ws.cell(r + 1, i + 1, v); cell.font = Font(italic=True)
        if i: cell.number_format = '0.00'
    ws.cell(r + 3, 1, "Notes").font = Font(bold=True)
    for i, n in enumerate([
        "Percentages recomputed as shares / total shares. Sub-category rows (indented) are shown for reference and are part of the bold bucket above them.",
        "KMP: SEBI format added this category from the Sep-2022 quarter; before that it is blank and sits inside Public.",
        "Pre-Sep-2022 filings reclassified for comparability: NBFCs -> DII, Overseas Corporate Bodies -> FII.",
        "Sep-21 is pre-bonus (29.96 crore shares); the 2:1 bonus in Dec-2021 took the count to 89.87 crore. The 08-Dec-2021 event filing is excluded here.",
    ]):
        ws.cell(r + 4 + i, 1, n)
    ws.column_dimensions["A"].width = 46
    for i in range(2, len(cols) + 2):
        ws.column_dimensions[get_column_letter(i)].width = 8.5
    ws.row_dimensions[2].height = 30
    ws.freeze_panes = "B3"

print(matrix.to_string(index=False))
print(f"\nwritten {path}")
